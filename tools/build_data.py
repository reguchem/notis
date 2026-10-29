"""엑셀 목록을 홈페이지 검색용 데이터로 변환합니다.
사용법: python tools/build_data.py   (저장소 폴더의 엑셀 파일을 자동으로 찾습니다)

- 이름에 '대리인'이 들어간 파일: 공동등록 완료 물질 목록 → data/substances.js
  (2행 머리글: 화학물질명, 고유번호, CAS No., 등록 예정시 문의처, 비고, 컨설팅사 / 4행부터 자료)
  ※ 첫 번째 시트만 사용합니다.
- 이름이 '별표'로 시작하는 파일: 기존화학물질 목록(별표1, 별표2) → data/existing.js
  (열: 고유번호, '화학물질명 (CAS No. ...)')
- 이름에 'CMR'이 들어간 PDF: 등록대상 CMR물질 고시(안) 별표 → data/cmr.js  (pdfplumber 필요)
"""
import glob, json, os, re
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
xlsx = [f for f in glob.glob(os.path.join(ROOT, '*.xlsx')) if not os.path.basename(f).startswith('~$')]
CAS_RE = re.compile(r'(\d{2,7})\s*-\s*(\d{2})\s*-\s*(\d)(?!\d)')

def s(v):
    return re.sub(r'\s+', ' ', str(v)).strip() if v is not None else ''

def asof_of(path):
    m = re.search(r'_(\d{6})', os.path.basename(path))
    return f"20{m.group(1)[:2]}.{m.group(1)[2:4]}.{m.group(1)[4:]}" if m else ''

def write(name, var, obj):
    out = os.path.join(ROOT, 'data', name)
    with open(out, 'w', encoding='utf-8') as f:
        f.write('// 자동 생성 파일: tools/build_data.py 로 다시 만드세요.\n')
        f.write(f'window.{var} = ' + json.dumps(obj, ensure_ascii=False, separators=(',', ':')) + ';\n')
    return out

# 1) 공동등록 완료 목록
joint = [f for f in xlsx if '대리인' in os.path.basename(f)]
if joint:
    src = sorted(joint)[-1]
    ws = openpyxl.load_workbook(src, read_only=True, data_only=True).worksheets[0]
    header = [s(c) for c in next(ws.iter_rows(min_row=2, max_row=2, values_only=True))]
    def find(name):  # 머리글이 '컨설팅사(또는 대표자)'처럼 조금 달라도 찾습니다
        for i, h in enumerate(header):
            if h.replace(' ', '').startswith(name.replace(' ', '')):
                return i
        raise SystemExit(f'엑셀 머리글에서 "{name}" 열을 찾지 못했습니다: {header}')
    col = {n: find(n) for n in ['화학물질명', '고유번호', 'CAS No.', '등록 예정시 문의처', '비고', '컨설팅사']}
    items = []
    for r in ws.iter_rows(min_row=4, values_only=True):
        name, ids, cas = s(r[col['화학물질명']]), s(r[col['고유번호']]), s(r[col['CAS No.']])
        if not (name or ids or cas):
            continue
        items.append({'n': name, 'id': [x.strip() for x in ids.split(',') if x.strip()],
                      'cas': ['-'.join(m) for m in CAS_RE.findall(cas)], 'casRaw': cas,
                      'tel': s(r[col['등록 예정시 문의처']]), 'note': s(r[col['비고']]), 'org': s(r[col['컨설팅사']])})
    write('substances.js', 'SUBSTANCE_DATA', {'asof': asof_of(src), 'source': os.path.basename(src), 'items': items})
    print(f'공동등록 완료 {len(items)}건 ← {os.path.basename(src)}')

# 2) 기존화학물질 목록 (별표1·2)
lists, total = [], 0
for src in sorted(f for f in xlsx if os.path.basename(f).startswith('별표')):
    label = '별표' + re.match(r'별표(\d+)', os.path.basename(src)).group(1)
    ws = openpyxl.load_workbook(src, read_only=True, data_only=True).worksheets[0]
    rows, removed = [], 0
    for r in ws.iter_rows(values_only=True):
        id_, text = s(r[0]), s(r[1] if len(r) > 1 else '')
        if not id_ or not text or id_ == '고유번호' or id_.startswith('[') or id_.startswith('*'):
            continue
        if text == '삭제':
            removed += 1
            continue
        i = text.rfind('CAS No')
        name = re.sub(r'[\s;(]+$', '', text[:i]).strip() if i >= 0 else text
        cas = ['-'.join(m) for m in CAS_RE.findall(text[i:])] if i >= 0 else []
        rows.append([id_, name, cas])
    lists.append({'label': label, 'file': os.path.basename(src), 'rows': rows})
    total += len(rows)
    print(f'{label} {len(rows)}건 (삭제 {removed}건 제외) ← {os.path.basename(src)}')
if lists:
    write('existing.js', 'EXISTING_DATA', {'lists': lists})
    print(f'기존화학물질 합계 {total}건')

# 3) 등록대상 CMR물질 고시(안) (PDF 표: 순번, 화학물질 명칭, 고유번호(CAS No.))
cmr_pdf = [f for f in glob.glob(os.path.join(ROOT, '*.pdf')) if 'CMR' in os.path.basename(f).upper()]
if cmr_pdf:
    try:
        import pdfplumber
    except ImportError:
        pdfplumber = None
        print('pdfplumber가 없어 CMR 목록은 건너뜁니다 (pip install pdfplumber)')
    if pdfplumber:
        src = sorted(cmr_pdf)[-1]
        rows = []
        with pdfplumber.open(src) as pdf:
            for page in pdf.pages:
                for t in page.extract_tables():
                    for r in t:
                        if r and r[0] and re.fullmatch(r'\d+', r[0].strip()):
                            name, cas = s(r[1]), s(r[2])
                            rows.append({'no': int(r[0]), 'n': name, 'cas': ['-'.join(m) for m in CAS_RE.findall(cas)], 'casRaw': cas})
        rows.sort(key=lambda x: x['no'])
        write('cmr.js', 'CMR_DATA', {'title': '‘21년까지 등록하여야 할 암, 돌연변이, 생식능력 이상을 일으키거나 일으킬 우려가 있는 기존화학물질',
                                     'notice': '환경부고시 제2018-232호(안)', 'source': os.path.basename(src), 'items': rows})
        print(f'등록대상 CMR물질 {len(rows)}건 ← {os.path.basename(src)}')

# 4) 견적 산출 (견적산출.xlsx의 '톤수별 견적', '현재환율' 시트) → data/quotes.js
#    공개 범위를 좁히기 위해 QUOTE_CAS에 적은 물질만 내보냅니다.
QUOTE_CAS = ['1317-38-0']
ITEMS = ['개별서류 작성 컨설팅', 'CSR 노출평가 컨설팅', '별지 26호 서식작성 컨설팅', '등록행정수수료', '협의체 참조권',
         'CSR 유해성평가 참조권', '시험자료 참조권', '정부생산자료 참조권', '기타 참조권']
qx = sorted(glob.glob(os.path.join(ROOT, '**', '견적산출*.xlsx'), recursive=True))
qx = [f for f in qx if not os.path.basename(f).startswith('~$')]
if qx:
    src = qx[-1]
    wb = openpyxl.load_workbook(src, data_only=True)
    ws = wb['톤수별 견적']
    rates = {}
    for r in wb['현재환율'].iter_rows(min_row=2, values_only=True):
        if r and r[0] and r[2]:
            rates[str(r[0]).strip()] = {'send': float(r[2]), 'date': r[4].strftime('%Y-%m-%d') if hasattr(r[4], 'strftime') else s(r[4])}
    subs = {}
    rows = list(ws.iter_rows(values_only=True))
    for i, r in enumerate(rows):
        if r[0] == '물질정보' and s(r[2]) in QUOTE_CAS:
            cas, name = s(r[2]), s(r[5])
            block = {}
            for r2 in rows[i + 3:i + 13]:
                if r2[0] in ITEMS:
                    vals = [v if isinstance(v, (int, float)) else None for v in r2[1:5]]
                    block[r2[0]] = {'v': vals, 'note': s(r2[5]), 'sel': s(r2[7])}
            subs[cas] = {'name': name, 'items': block, 'testCurrency': block.get('시험자료 참조권', {}).get('sel') or 'KRW',
                         'carcinogen': block.get('CSR 노출평가 컨설팅', {}).get('sel')}
    write('quotes.js', 'QUOTE_DATA', {'source': os.path.basename(src), 'tonnage': ['1-10톤', '10-100톤', '100-1000톤', '1000톤 이상'],
                                      'items': ITEMS, 'rates': rates, 'substances': subs})
    print(f'견적 산출 {len(subs)}종 ← {os.path.basename(src)} (환율 {", ".join(f"{k} {v[chr(115)+chr(101)+chr(110)+chr(100)]}" for k, v in rates.items())})')

# 5) 고분자화합물로서 등록면제 확인받은 기존화학물질 목록 (열: 고유번호, CAS번호, 물질명) → data/polymers.js
px = [f for f in glob.glob(os.path.join(ROOT, '*.xlsx')) if '고분자' in os.path.basename(f) and not os.path.basename(f).startswith('~$')]
if px:
    src = sorted(px)[-1]
    ws = openpyxl.load_workbook(src, read_only=True, data_only=True).worksheets[0]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not any(r):
            continue
        rows.append([s(r[0]), ['-'.join(m) for m in CAS_RE.findall(s(r[1]))], s(r[2])])
    write('polymers.js', 'POLYMER_DATA', {'source': os.path.basename(src), 'items': rows})
    print(f'고분자 등록면제 확인 {len(rows)}건 ← {os.path.basename(src)}')
